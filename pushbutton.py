import bs4
import pandas
import openpyxl
import requests




class FiEx:

    def __init__(self, url: str):
        self.url = url

    def htmlDocument(self) -> bs4.BeautifulSoup:
        with open(self.url, 'r') as mainURL:
            body = mainURL.read()
            soup = bs4.BeautifulSoup(body, 'html.parser')
        return soup

    def req_htmlDocument(self):
        resp = requests.get(self.url)
        soup = bs4.BeautifulSoup(resp.text, 'html.parser')
        return soup


class ExtractData:

    def __init__(self, body: bs4.BeautifulSoup):
        self.body = body

    def exData(self) -> list:
        data_table = self.body.find(id='modality')
        sata = []
        for tr in data_table.find_all('tr'):
            mata = []
            for td in tr.find_all('td'):
                mata.append(td.text)
            sata.append(mata)
        return sata

    def ex_shareholding_details(self):
        soup = self.body
        trs = soup.find(class_='shareholdcnt')
        lst = []
        for tr in trs.find['tr']:
            lst1 = []
            for td in tr.find['td']:
                lst1.append[td.text]
            lst.append[lst]
        return lst





    #     dataframe.set_index('Stokes')
    #     dataframe.rename(columns={'Unnamed: 43': 'URL'}, inplace=True)
    #     dataframe = dataframe[['Stokes', 'URL']]
    #     dataframe.dropna(axis='index', how='any', inplace=True)
    #     dataframe = dataframe.replace('https://www.financialexpress.com', 'NA')
    #     dataframe = dataframe[dataframe['URL'] != 'NA']
    #     dataframe = dataframe.set_index('Stokes')
    #     stocke_URL = dataframe.loc['Mangalam Timber Prod', 'URL']
    #     return Mangalam


if __name__=='__main__':
    ob = FiEx("C:\\Users\\DELL\\Desktop\\jk.txt")
    d = ob.htmlDocument()
    print(d.text)
    p = ExtractData(d).ex_shareholding_details()
    print(p)




